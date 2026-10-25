#!/usr/bin/env python3
"""Chuyển bộ tài liệu Markdown của repo thành một tệp .docx in được.

Dùng: python3 tools/md2docx.py
Yêu cầu: python-docx (pip install python-docx)
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "DE-CUONG-ON-TAP-NGU-VAN-10-GK1.docx"

FILES = [
    ("PHẦN 1", "01-PHAN-TICH-MA-TRAN-VA-CHIEN-LUOC-ON.md"),
    ("PHẦN 2", "02-LY-THUYET-THE-LOAI-THAN-THOAI.md"),
    ("PHẦN 3", "03-CAC-VAN-BAN-TRONG-SGK.md"),
    ("PHẦN 4", "04-KY-NANG-DOC-HIEU.md"),
    ("PHẦN 5", "05-DE-KIEM-TRA-VA-DAP-AN.md"),
    ("PHẦN 6", "06-KY-NANG-VIET-NGHI-LUAN.md"),
    ("PHẦN 7", "07-SO-DO-TU-DUY-MEO-GHI-NHO.md"),
    ("PHẦN 8", "08-DE-LUYEN-TAP-BO-SUNG.md"),
    ("PHẦN 9", "09-PHU-LUC-NGU-LIEU.md"),
]

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"


def set_style(doc):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(12.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.25
    for name, size, color in (
        ("Heading 1", 20, (0x0B, 0x3D, 0x91)),
        ("Heading 2", 15, (0x0B, 0x3D, 0x91)),
        ("Heading 3", 13, (0x1F, 0x4E, 0x79)),
    ):
        s = doc.styles[name]
        s.font.name = BODY_FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(*color)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)


def add_inline(par, text, base_bold=False, base_italic=False, mono=False):
    """Ghi text có **đậm**, *nghiêng*, `code` vào paragraph."""
    tokens = re.split(r"(\*\*.+?\*\*|`.+?`|(?<!\*)\*[^*\n]+?\*(?!\*))", text)
    for tok in tokens:
        if not tok:
            continue
        bold, italic, code = base_bold, base_italic, mono
        content = tok
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            content, bold = tok[2:-2], True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            content, code = tok[1:-1], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            content, italic = tok[1:-1], True
        run = par.add_run(content)
        run.bold = bold
        run.italic = italic
        run.font.name = MONO_FONT if code else BODY_FONT
        if code:
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x7A, 0x2A, 0x00)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], cell_text, base_bold=True)
    for row in body:
        cells = table.add_row().cells
        for i in range(ncols):
            txt = row[i] if i < len(row) else ""
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], txt)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                for r in p.runs:
                    r.font.size = Pt(10.5)
    doc.add_paragraph()


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def convert(doc, md_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, in_code, code_buf = 0, False, []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # code fence
        if stripped.startswith("```"):
            if in_code:
                for idx, codeline in enumerate(code_buf):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    if idx == 0:
                        p.paragraph_format.space_before = Pt(4)
                    if idx == len(code_buf) - 1:
                        p.paragraph_format.space_after = Pt(8)
                    run = p.add_run(codeline if codeline else " ")
                    run.font.name = MONO_FONT
                    run.font.size = Pt(9)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\s*\|?[\s:|-]+\|[\s:|-]*$", lines[i + 1]
        ):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                doc.add_page_break()
            p = doc.add_heading(level=min(level, 3))
            add_inline(p, text)
            i += 1
            continue

        # horizontal rule
        if re.match(r"^-{3,}$", stripped) or re.match(r"^\*{3,}$", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, " ".join(x for x in buf if x), base_italic=False)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        # list items
        m = re.match(r"^[-*]\s+\[( |x|X)\]\s+(.*)$", stripped)
        if m:
            box = "☐ " if m.group(1) == " " else "☑ "
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, box + m.group(2))
            i += 1
            continue
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
            i += 1
            continue
        m = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, f"{m.group(1)}. {m.group(2)}")
            i += 1
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, stripped)
        i += 1


def title_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ĐỀ CƯƠNG ÔN TẬP GIỮA HỌC KÌ I")
    r.bold = True
    r.font.size = Pt(26)
    r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor(0x0B, 0x3D, 0x91)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NGỮ VĂN 10 — BỘ SÁCH KẾT NỐI TRI THỨC VỚI CUỘC SỐNG")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Năm học 2026 – 2027")
    r.font.size = Pt(13)

    doc.add_paragraph()
    info = [
        ("Trọng tâm", "Truyện thần thoại (ngữ liệu ngoài SGK) + Viết văn bản nghị luận phân tích, đánh giá tác phẩm truyện"),
        ("Cấu trúc đề", "11 câu — 8 trắc nghiệm + 2 tự luận đọc hiểu + 1 bài viết"),
        ("Tỉ lệ điểm", "Đọc hiểu 60% — Viết 40%"),
        ("Thời gian", "90 phút"),
        ("Nguồn biên soạn", "Ma trận & bản đặc tả đề kiểm tra giữa kì 1 môn Ngữ văn 10, năm học 2026 – 2027"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in info:
        cells = table.add_row().cells
        cells[0].text = ""
        add_inline(cells[0].paragraphs[0], k, base_bold=True)
        cells[1].text = ""
        add_inline(cells[1].paragraphs[0], v)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        "Bộ tài liệu gồm 9 phần:",
        "lý thuyết thể loại · văn bản SGK · kỹ năng đọc hiểu · đề kiểm tra đầy đủ kèm đáp án",
        "kỹ năng viết · sơ đồ & mẹo nhớ · đề luyện tập · phụ lục ngữ liệu",
    ]
    for k, line in enumerate(lines):
        pp = p if k == 0 else doc.add_paragraph()
        if k:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(line)
        rr.italic = True
        rr.font.size = Pt(11)
    doc.add_page_break()


def main():
    doc = Document()
    set_style(doc)
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(1.8)

    # header/footer
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Đề cương ôn tập Giữa kì 1 – Ngữ văn 10 (KNTT) – 2026–2027")
    run.font.size = Pt(9)
    run.font.name = BODY_FONT

    title_page(doc)
    for _, fname in FILES:
        path = ROOT / fname
        if not path.exists():
            print(f"[!] thiếu {fname}", file=sys.stderr)
            continue
        convert(doc, path)

    doc.save(OUT)
    print(f"Đã tạo {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
