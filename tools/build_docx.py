#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ghép 6 tệp Markdown trong docs/ thành một tệp .docx hoàn chỉnh
(đề cương ôn tập giữa kì 1 Ngữ văn 10 — Kết nối tri thức).

Cách dùng:
    python3 tools/build_docx.py

Yêu cầu: pip install python-docx
Kết quả: DE-CUONG-ON-TAP-GIUA-KI-1-VAN-10.docx (ở thư mục gốc repo)
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")
OUT = os.path.join(ROOT, "DE-CUONG-ON-TAP-GIUA-KI-1-VAN-10.docx")

BODY_FONT = "Times New Roman"
HEAD_FONT = "Times New Roman"
BODY_SIZE = Pt(13)
ACCENT = RGBColor(0x1F, 0x3B, 0x57)  # xanh đen
ACCENT2 = RGBColor(0x8B, 0x1A, 0x1A)  # đỏ nâu

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


# --------------------------------------------------------------------------- #
# Tiện ích định dạng
# --------------------------------------------------------------------------- #
def set_run_font(run, size=BODY_SIZE, bold=False, italic=False, color=None,
                 name=BODY_FONT):
    run.font.name = name
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def add_inline(paragraph, text, size=BODY_SIZE, base_bold=False, color=None):
    """Thêm text có **đậm**, `code`, *nghiêng* vào paragraph."""
    text = text.replace("\\|", "|")
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            set_run_font(run, size, bold=True, color=color)
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            set_run_font(run, size, bold=base_bold)
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1].replace("`", ""))
            set_run_font(run, size, bold=base_bold, italic=True, color=color)
        else:
            run = paragraph.add_run(chunk)
            set_run_font(run, size, bold=base_bold, color=color)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_hr(document):
    p = document.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B0B0B0")
    borders.append(bottom)
    p_pr.append(borders)


# --------------------------------------------------------------------------- #
# Bộ phân tích Markdown tối giản
# --------------------------------------------------------------------------- #
def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator_row(cells):
    return all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells if c) and any(cells)


def flush_paragraph(document, buffer):
    if not buffer:
        return
    text = " ".join(buffer).strip()
    buffer.clear()
    if not text:
        return
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.25
    add_inline(p, text)


def add_table(document, rows):
    if not rows:
        return
    header, body = rows[0], rows[1:]
    ncols = max(len(r) for r in rows)
    table = document.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i in range(ncols):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value = header[i] if i < len(header) else ""
        add_inline(p, value, size=Pt(11), base_bold=True)
        for run in p.runs:
            set_run_font(run, Pt(11), bold=True)
        shade(cell, "DCE6F1")

    for row in body:
        cells = table.add_row().cells
        for i in range(ncols):
            value = row[i] if i < len(row) else ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, size=Pt(11))
            for run in p.runs:
                set_run_font(run, Pt(11))

    document.add_paragraph().paragraph_format.space_after = Pt(2)


def convert_file(document, path):
    with open(path, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    buffer = []
    table_rows = []
    in_code = False
    code_buffer = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            rows = [r for r in table_rows if not is_separator_row(r)]
            add_table(document, rows)
        table_rows = []

    for raw in lines:
        line = raw.rstrip()

        # khối code
        if line.strip().startswith("```"):
            flush_paragraph(document, buffer)
            flush_table()
            if in_code:
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                run = p.add_run("\n".join(code_buffer))
                set_run_font(run, Pt(10), name="Consolas")
                code_buffer.clear()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buffer.append(raw)
            continue

        # bảng
        if line.strip().startswith("|"):
            flush_paragraph(document, buffer)
            table_rows.append(split_row(line))
            continue
        else:
            flush_table()

        # tiêu đề
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush_paragraph(document, buffer)
            level, text = len(m.group(1)), m.group(2).strip()
            if level == 1:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(10)
                add_inline(p, text, size=Pt(19), base_bold=True, color=ACCENT)
            elif level == 2:
                p = document.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                add_inline(p, text, size=Pt(15), base_bold=True, color=ACCENT2)
                p_pr = p._p.get_or_add_pPr()
                borders = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "8")
                bottom.set(qn("w:space"), "2")
                bottom.set(qn("w:color"), "8B1A1A")
                borders.append(bottom)
                p_pr.append(borders)
            else:
                p = document.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                size = Pt(13.5) if level == 3 else Pt(13)
                add_inline(p, text, size=size, base_bold=True, color=ACCENT)
            continue

        # đường kẻ
        if re.fullmatch(r"\s*(-{3,}|\*{3,}|_{3,})\s*", line):
            flush_paragraph(document, buffer)
            add_hr(document)
            continue

        # trích dẫn
        if line.strip().startswith(">"):
            flush_paragraph(document, buffer)
            text = line.strip().lstrip(">").strip()
            if not text:
                continue
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            p_pr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), "1F3B57")
            borders.append(left)
            p_pr.append(borders)
            add_inline(p, text, size=Pt(12.5))
            for run in p.runs:
                if not run.bold:
                    run.italic = True
            continue

        # checklist
        m = re.match(r"^\s*-\s*\[( |x|X)\]\s*(.*)$", line)
        if m:
            flush_paragraph(document, buffer)
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.space_after = Pt(2)
            box = "☒ " if m.group(1).lower() == "x" else "☐ "
            run = p.add_run(box)
            set_run_font(run, BODY_SIZE)
            add_inline(p, m.group(2))
            continue

        # danh sách
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m:
            flush_paragraph(document, buffer)
            indent = Cm(0.6 + 0.5 * (len(m.group(1)) // 2))
            p = document.add_paragraph()
            p.paragraph_format.left_indent = indent
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("• ")
            set_run_font(run, BODY_SIZE)
            add_inline(p, m.group(2))
            continue

        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if m:
            flush_paragraph(document, buffer)
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 + 0.5 * (len(m.group(1)) // 2))
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(m.group(2) + ". ")
            set_run_font(run, BODY_SIZE, bold=True)
            add_inline(p, m.group(3))
            continue

        if not line.strip():
            flush_paragraph(document, buffer)
            continue

        buffer.append(line.strip())

    flush_paragraph(document, buffer)
    flush_table()
    if code_buffer:
        p = document.add_paragraph()
        run = p.add_run("\n".join(code_buffer))
        set_run_font(run, Pt(10), name="Consolas")


# --------------------------------------------------------------------------- #
def build():
    doc = Document()

    # Thiết lập trang A4
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.8)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    # Trang bìa
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "ĐỀ CƯƠNG ÔN TẬP", size=Pt(26), base_bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "GIỮA HỌC KÌ I", size=Pt(30), base_bold=True, color=ACCENT2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "NGỮ VĂN 10", size=Pt(24), base_bold=True, color=ACCENT)
    doc.add_paragraph()
    for line in [
        "Bộ sách Kết nối tri thức với cuộc sống",
        "Năm học 2026 – 2027",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, line, size=Pt(14), base_bold=True)
    doc.add_paragraph()
    for line in [
        "Biên soạn bám sát MA TRẬN KIỂM TRA GIỮA KÌ 1",
        "Cấu trúc đề: 11 câu / 10 điểm / 90 phút",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, line, size=Pt(12))
    doc.add_paragraph()
    for line in [
        "6 phần · 6 đề đầy đủ · 66 câu hỏi có đáp án",
        "3 bài văn mẫu · 10 dàn ý chi tiết",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, line, size=Pt(12), base_bold=True)
    doc.add_page_break()

    files = sorted(f for f in os.listdir(DOCS) if f.endswith(".md"))
    if not files:
        sys.exit("Không tìm thấy tệp .md nào trong docs/")

    for idx, name in enumerate(files):
        if idx:
            doc.add_page_break()
        convert_file(doc, os.path.join(DOCS, name))

    # Đánh số trang
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    set_run_font(run, Pt(10))

    doc.save(OUT)
    print("Đã tạo:", OUT)
    print("Số tệp nguồn:", len(files))
    print("Số đoạn văn:", len(doc.paragraphs))
    print("Số bảng:", len(doc.tables))


if __name__ == "__main__":
    build()
