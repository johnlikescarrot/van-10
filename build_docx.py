#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Dựng bản .docx in được từ các file Markdown trong thư mục de-cuong/.
Chạy:  python3 build_docx.py
Kết quả: output/De-cuong-on-tap-Giua-ki-1-Ngu-van-10.docx
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "de-cuong")
OUT_DIR = os.path.join(BASE, "output")
OUT = os.path.join(OUT_DIR, "De-cuong-on-tap-Giua-ki-1-Ngu-van-10.docx")

FILES = [
    ("01-the-loai-than-thoai-va-truyen-ke.md", "PHẦN I"),
    ("02-ky-nang-doc-hieu-va-dang-cau-hoi.md", "PHẦN II"),
    ("03-thuc-hanh-tieng-viet-tu-han-viet.md", "PHẦN III"),
    ("04-ky-nang-viet-nghi-luan.md", "PHẦN IV"),
    ("05-de-kiem-tra-chuan-ma-tran-va-dap-an.md", "PHẦN V"),
    ("06-bo-de-luyen-tap.md", "PHẦN VI"),
    ("07-so-tay-ghi-nho-va-lich-on.md", "PHẦN VII"),
]

FONT = "Times New Roman"
BODY_SIZE = 12.5

# ---------------------------------------------------------------- helpers


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def style_document(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_SIZE)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.25

    for lvl, size, color in ((1, 17, "1F3864"), (2, 14.5, "2E5496"),
                             (3, 13, "2E5496"), (4, 12.5, "333333")):
        s = doc.styles["Heading %d" % lvl]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = RGBColor.from_string(color)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.paragraph_format.space_before = Pt(12 if lvl == 1 else 9)
        s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.keep_with_next = True


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_inline(par, text, base_bold=False, base_italic=False, size=None, color=None):
    """Xử lý **bold**, *italic*, `code` trong một dòng (hỗ trợ lồng nhau)."""
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]*?`|\*[^*\n]+?\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_inline(par, tok[2:-2], True, base_italic, size, color)
            continue
        if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            run.bold = base_bold
            run.italic = base_italic
            run.font.name = "Consolas"
            if size:
                run.font.size = Pt(size)
            continue
        if tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            add_inline(par, tok[1:-1], base_bold, True, size, color)
            continue
        run = par.add_run(tok)
        run.bold = base_bold
        run.italic = base_italic
        run.font.name = FONT
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, rows):
    header = rows[0]
    body = rows[1:]
    ncol = len(header)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    hdr = table.rows[0].cells
    for i, cell_text in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        add_inline(p, cell_text, base_bold=True, size=10.5)
        set_cell_bg(hdr[i], "DCE6F1")

    for r in body:
        cells = table.add_row().cells
        for i in range(ncol):
            txt = r[i] if i < len(r) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, txt, size=10.5)

    # repeat header row
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.4)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("1F3864")


def add_quote(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.7)
        pf.space_after = Pt(2)
        add_inline(p, line, base_italic=True, size=11.5)


# ---------------------------------------------------------------- parser


def parse_table(lines, i):
    """Đọc một bảng markdown bắt đầu ở dòng i. Trả về (rows, next_index)."""
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip()
        cells = [c.strip() for c in raw.strip("|").split("|")]
        rows.append(cells)
        i += 1
    # bỏ dòng phân cách |---|---|
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c)]
    return rows, i


def convert_file(doc, path, first=False):
    with open(path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    if not first:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        # code fence
        if s.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1
            add_code_block(doc, buf)
            continue

        # blank
        if not s:
            i += 1
            continue

        # table
        if s.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                add_table(doc, rows)
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}", s):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "BFBFBF")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            level = min(len(m.group(1)), 4)
            hp = doc.add_paragraph(style="Heading %d" % level)
            add_inline(hp, m.group(2).strip())
            i += 1
            continue

        # quote
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            add_quote(doc, buf)
            continue

        # bullet list
        if re.match(r"^[-*]\s+", s):
            txt = re.sub(r"^[-*]\s+", "", s)
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.75)
            add_inline(p, txt)
            i += 1
            continue

        # ordered list
        m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.75)
            add_inline(p, m.group(2))
            i += 1
            continue

        # paragraph (gộp các dòng liền nhau)
        buf = [s]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "```"))
                    or re.match(r"^[-*]\s+", nxt)
                    or re.match(r"^\d+[.)]\s+", nxt)
                    or re.fullmatch(r"-{3,}", nxt)):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, " ".join(buf))


# ---------------------------------------------------------------- main


def build_cover(doc):
    for _ in range(3):
        doc.add_paragraph()

    def center(text, size, bold=False, color=None, space=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = RGBColor.from_string(color)

    center("ĐỀ CƯƠNG ÔN TẬP", 15, True, "333333")
    center("GIỮA HỌC KÌ I", 30, True, "1F3864", 2)
    center("MÔN NGỮ VĂN — LỚP 10", 20, True, "2E5496", 18)
    center("Bộ sách: KẾT NỐI TRI THỨC VỚI CUỘC SỐNG", 13, False, "333333", 4)
    center("Năm học 2026 – 2027", 13, False, "333333", 24)
    center("Biên soạn bám sát MA TRẬN và BẢN ĐẶC TẢ đề kiểm tra của nhà trường", 11.5, False, "555555", 4)
    center("Trọng tâm: TRUYỆN THẦN THOẠI — Đọc hiểu & Viết văn nghị luận", 11.5, False, "555555", 30)

    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    cell = box.rows[0].cells[0]
    set_cell_bg(cell, "F2F5FA")
    cell.text = ""
    items = [
        "Phần I      Thể loại thần thoại và truyện kể",
        "Phần II     Kĩ năng đọc hiểu: 12 dạng câu hỏi và công thức trả lời",
        "Phần III    Thực hành tiếng Việt: sử dụng từ Hán Việt",
        "Phần IV     Kĩ năng viết nghị luận phân tích, đánh giá tác phẩm truyện",
        "Phần V      Đề kiểm tra chuẩn ma trận và đáp án chi tiết",
        "Phần VI     Bộ đề luyện tập (3 đề + đáp án)",
        "Phần VII    Sổ tay ghi nhớ, flashcard và lịch ôn tập 7 ngày",
    ]
    first = True
    for it in items:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(it)
        r.font.name = FONT
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor.from_string("1F3864")

    doc.add_paragraph()
    center("Tài liệu gồm 7 phần — dùng được cho cả ôn thi và dạy học", 10.5, False, "777777")


def build_toc(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("MỤC LỤC", level=1)
    toc = [
        ("PHẦN I — THỂ LOẠI THẦN THOẠI VÀ TRUYỆN KỂ", ""),
        ("", "0. Bản đồ phạm vi ôn thi"),
        ("", "1. Truyện kể — sáu khái niệm gốc"),
        ("", "2. Thần thoại — trái tim của đề thi"),
        ("", "3. Hệ thống văn bản trong SGK Bài 1"),
        ("", "4. Mở rộng phạm vi — vượt ra ngoài SGK"),
        ("", "5. Bảng ghi nhớ nhanh + mnemonic"),
        ("PHẦN II — KĨ NĂNG ĐỌC HIỂU: 12 DẠNG CÂU HỎI", ""),
        ("", "1. Giải mã phần đọc hiểu của đề"),
        ("", "2. Quy trình 7 bước làm bài"),
        ("", "3. 12 dạng câu hỏi và khung trả lời"),
        ("", "4. Bộ câu mở đoạn dùng lại được"),
        ("", "5. Mười lỗi mất điểm oan"),
        ("", "6. Tự luyện nhanh + đáp án"),
        ("PHẦN III — THỰC HÀNH TIẾNG VIỆT: TỪ HÁN VIỆT", ""),
        ("", "1. Từ Hán Việt là gì?"),
        ("", "2. Phân loại từ Hán Việt"),
        ("", "3. Vai trò và sắc thái"),
        ("", "4. Sử dụng từ Hán Việt đúng cách"),
        ("", "5. Từ Hán Việt trong các văn bản Bài 1"),
        ("", "6. Bài tập chuẩn theo SGK + đáp án"),
        ("", "7. Từ Hán Việt trong bài viết nghị luận"),
        ("", "8. Bảng 60 yếu tố Hán Việt gốc"),
        ("", "9. Trắc nghiệm tự luyện + đáp án"),
        ("PHẦN IV — KĨ NĂNG VIẾT NGHỊ LUẬN", ""),
        ("", "1. Yêu cầu của đề — đọc từ bản đặc tả"),
        ("", "2. Thang điểm 4,0 và cách ăn từng phần"),
        ("", "3. Dàn ý tổng quát — bộ khung vạn năng"),
        ("", "4. Ngân hàng luận điểm — 7 nhóm đặc sắc nghệ thuật"),
        ("", "5. Máy đào luận điểm — 12 câu hỏi tự vấn"),
        ("", "6. Bài viết mẫu 1 (700 chữ, có chú thích)"),
        ("", "7. Bài viết mẫu 2 (500 chữ, an toàn)"),
        ("", "8. Kho dẫn chứng dùng lại được"),
        ("", "9. Bộ từ ngữ đánh giá"),
        ("", "10. Mười lỗi khiến bài viết mất điểm"),
        ("", "11. Nói và nghe"),
        ("", "12. Tóm tắt 1 trang"),
        ("PHẦN V — ĐỀ KIỂM TRA CHUẨN MA TRẬN VÀ ĐÁP ÁN", ""),
        ("", "Ghi chú quan trọng về nguồn đề"),
        ("", "Phần A — Ma trận đối chiếu từng câu hỏi"),
        ("", "Đề kiểm tra giữa học kì I (11 câu, 90 phút)"),
        ("", "Đáp án và thang điểm chi tiết"),
        ("", "Rubric chấm câu Viết"),
        ("", "Phần B — Bảng tự chấm sau khi làm bài"),
        ("PHẦN VI — BỘ ĐỀ LUYỆN TẬP", ""),
        ("", "Đề số 2 + đáp án"),
        ("", "Đề số 3 + đáp án"),
        ("", "Đề số 4 (đề mở) + đáp án"),
        ("PHẦN VII — SỔ TAY GHI NHỚ VÀ LỊCH ÔN TẬP", ""),
        ("", "1. Lịch ôn tập 7 ngày"),
        ("", "2. Sổ tay ghi nhớ — 6 trang in dán"),
        ("", "3. Flashcard — 40 thẻ"),
        ("", "4. Sơ đồ tư duy tóm tắt"),
        ("", "5. Tam giác vàng — 3 điều quyết định điểm"),
        ("", "6. Checklist trước khi vào phòng thi"),
        ("", "7. Phụ lục A — nếu đề ra vào Bài 2 (thơ ca)"),
        ("", "8. Phụ lục B — 10 câu tự vấn cuối cùng"),
    ]
    for main, sub in toc:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if main:
            r = p.add_run(main)
            r.bold = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor.from_string("1F3864")
            r.font.name = FONT
            p.paragraph_format.space_before = Pt(8)
        else:
            p.paragraph_format.left_indent = Cm(0.8)
            r = p.add_run(sub)
            r.font.size = Pt(11)
            r.font.name = FONT
            r.font.color.rgb = RGBColor.from_string("333333")


def main():
    doc = Document()
    style_document(doc)

    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(1.8)
    add_page_number_footer(sec)

    build_cover(doc)
    build_toc(doc)

    for idx, (fname, _) in enumerate(FILES):
        path = os.path.join(SRC, fname)
        if not os.path.exists(path):
            print("  ! thiếu file:", fname)
            continue
        convert_file(doc, path, first=False)
        print("  + đã thêm:", fname)

    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(OUT)
    print("\nĐã tạo:", OUT, "(%.1f KB)" % (os.path.getsize(OUT) / 1024))


if __name__ == "__main__":
    main()
